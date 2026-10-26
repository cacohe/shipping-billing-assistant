import json
from decimal import Decimal
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from openpyxl import Workbook

from adapters.outbound.llm.langchain_extractor import LangChainBillingExtractor
from adapters.outbound.parsers.composite import CompositeDocumentParser
from application.dto.documents import InputFile
from application.ports.outbound.ocr import OcrResult

FIXTURES = Path(__file__).parents[2] / "fixtures" / "documents"


class Ocr:
    def recognize(self, file_url: str) -> OcrResult:
        return OcrResult("Invoice 1", (("charge", "amount"), ("freight", 10)), {"url": file_url})


class NoChargeOcr:
    def recognize(self, file_url: str) -> OcrResult:
        return OcrResult("FREIGHT PREPAID", (), {"url": file_url})


class Store:
    def __init__(self) -> None:
        self.uploaded: list[str] = []
        self.deleted: list[str] = []

    def upload(self, file: InputFile) -> str:
        self.uploaded.append(file.name)
        return f"ocr-inputs/{file.name}"

    def create_download_url(self, key: str, expires_in_seconds: int) -> str:
        assert expires_in_seconds == 600
        return f"https://storage.test/{key}"

    def delete(self, key: str) -> None:
        self.deleted.append(key)


@pytest.mark.integration
def test_xlsx_docx_eml_and_ocr_parsing() -> None:
    parser = CompositeDocumentParser(Ocr())
    workbook = Workbook()
    workbook.active.append(["charge", "amount"])
    workbook.active.append(["freight", 10])
    excel_bytes = BytesIO()
    workbook.save(excel_bytes)
    assert (
        parser.parse_one(InputFile("bill.xlsx", excel_bytes.getvalue())).tables[0].rows[0][1] == 10
    )

    word = Document()
    word.add_paragraph("Invoice 1")
    word_bytes = BytesIO()
    word.save(word_bytes)
    assert "Invoice" in parser.parse_one(InputFile("bill.docx", word_bytes.getvalue())).raw_text

    message = EmailMessage()
    message["Subject"] = "Bill"
    message.set_content("Please process")
    message.add_attachment(
        excel_bytes.getvalue(), maintype="application", subtype="octet-stream", filename="bill.xlsx"
    )
    parsed_email = parser.parse_one(InputFile("mail.eml", message.as_bytes()))
    assert parsed_email.children[0].source_type == "EXCEL"

    parsed_pdf = parser.parse_one(
        InputFile("bill.pdf", b"pdf", source_url="https://example.test/bill.pdf")
    )
    assert parsed_pdf.raw_text == "Invoice 1"


@pytest.mark.integration
def test_eml_pdf_attachment_uses_temporary_storage() -> None:
    store = Store()
    parser = CompositeDocumentParser(Ocr(), store)
    message = EmailMessage()
    message.set_content("Attached invoice")
    message.add_attachment(
        b"pdf",
        maintype="application",
        subtype="pdf",
        filename="invoice.pdf",
    )
    parsed = parser.parse_one(InputFile("mail.eml", message.as_bytes()))
    assert parsed.children[0].source_type == "PDF"
    assert parsed.children[0].raw_text == "Invoice 1"
    assert store.uploaded == ["invoice.pdf"]
    assert store.deleted == ["ocr-inputs/invoice.pdf"]


@pytest.mark.integration
def test_synthetic_samples_are_valid_parser_inputs() -> None:
    store = Store()
    parser = CompositeDocumentParser(Ocr(), store)
    email_path = next(FIXTURES.glob("*.eml"))
    expected_attachments = {
        "DEMO-MBL-20261026.pdf": "https://example.test/DEMO-MBL-20261026.pdf",
        "DEMO-HBL-20261026.pdf": "https://example.test/DEMO-HBL-20261026.pdf",
    }

    parsed_email = parser.parse_one(
        InputFile(
            email_path.name,
            email_path.read_bytes(),
            "message/rfc822",
        )
    )

    assert parsed_email.source_type == "EML"
    assert "DEMO-PREALERT-20261026" in parsed_email.metadata["subject"]
    assert {child.source_name for child in parsed_email.children} == set(expected_attachments)
    assert all(child.source_type == "PDF" for child in parsed_email.children)
    assert parsed_email.warnings == ("已忽略邮件正文内嵌资源：synthetic-logo.png",)

    for path in sorted((*FIXTURES.glob("*.pdf"), *FIXTURES.glob("*.png"))):
        parsed = parser.parse_one(
            InputFile(
                path.name,
                path.read_bytes(),
            )
        )
        expected_type = "PDF" if path.suffix.lower() == ".pdf" else "IMAGE"
        assert parsed.source_type == expected_type
        assert parsed.raw_payload == {"url": f"https://storage.test/ocr-inputs/{path.name}"}

    assert store.uploaded == [
        *expected_attachments,
        "synthetic-arrival-notice.pdf",
        "synthetic-debit-note.pdf",
        "synthetic-invoice.png",
    ]
    assert store.deleted == [f"ocr-inputs/{name}" for name in store.uploaded]


@pytest.mark.integration
def test_synthetic_eml_empty_envelope_recheck_extracts_only_grounded_charge() -> None:
    email_path = next(FIXTURES.glob("*.eml"))
    parsed_email = CompositeDocumentParser(NoChargeOcr(), Store()).parse_one(
        InputFile(email_path.name, email_path.read_bytes(), "message/rfc822")
    )
    reviewed_response = json.dumps(
        {
            "documents": [
                {
                    "document_info": {
                        "document_type": "PRE-ALERT",
                        "reference_no": None,
                        "document_date": "2026-10-26",
                        "bill_of_lading_no": "HBL-DEMO-20261026",
                    },
                    "parties": {"issuer": None, "recipient": None},
                    "charges": [
                        {
                            "raw_name": "H/C",
                            "amount": 50,
                            "currency": "USD",
                            "confidence": 0.8,
                            "evidence": [
                                {
                                    "source_id": "0",
                                    "source_name": email_path.name,
                                    "source_text": "Handling charge: H/C USD 50.00",
                                }
                            ],
                        }
                    ],
                }
            ]
        },
        ensure_ascii=False,
    )
    extractor = LangChainBillingExtractor(
        FakeListChatModel(responses=['{"documents": []}', reviewed_response])
    )

    result = extractor.extract((parsed_email,))

    assert len(result) == 1
    assert result[0].document_info.reference_no is None
    assert result[0].document_info.bill_of_lading_no == "HBL-DEMO-20261026"
    assert result[0].parties.issuer is None
    assert result[0].parties.recipient is None
    assert result[0].charge_direction is None
    assert len(result[0].charges) == 1
    charge = result[0].charges[0]
    assert charge.raw_name == "H/C"
    assert charge.amount == Decimal("50")
    assert charge.currency == "USD"
    assert charge.evidence[0].source_name == email_path.name
    assert charge.evidence[0].source_text == "Handling charge: H/C USD 50.00"
