"""DOCX parsing and optional legacy DOC conversion."""

import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document

from adapters.outbound.parsers.base import DocumentParseError
from application.dto.documents import (
    InputFile,
    ParsedDocument,
    SourceLocation,
    TableBlock,
    TextBlock,
)


class WordParser:
    def supports(self, file: InputFile) -> bool:
        return file.name.lower().endswith(".docx")

    def parse(self, file: InputFile) -> ParsedDocument:
        try:
            document = Document(BytesIO(file.content))
        except Exception as exc:
            raise DocumentParseError(f"Word 文件解析失败：{file.name}") from exc

        location = SourceLocation(source_name=file.name)
        text_blocks = tuple(
            TextBlock(paragraph.text.strip(), location)
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        tables: list[TableBlock] = []
        for table_index, table in enumerate(document.tables, start=1):
            rows = tuple(tuple(cell.text.strip() for cell in row.cells) for row in table.rows)
            if not rows:
                continue
            tables.append(
                TableBlock(
                    headers=rows[0],
                    rows=rows[1:],
                    location=SourceLocation(source_name=file.name, row=table_index),
                )
            )
        raw_text = "\n".join(block.text for block in text_blocks)
        return ParsedDocument(
            source_name=file.name,
            source_type="WORD",
            text_blocks=text_blocks,
            tables=tuple(tables),
            raw_text=raw_text,
        )


class LegacyDocParser:
    """Convert legacy binary DOC through LibreOffice, then reuse the DOCX parser."""

    def __init__(self, command: str = "soffice") -> None:
        self._command = command
        self._docx_parser = WordParser()

    def supports(self, file: InputFile) -> bool:
        return file.name.lower().endswith(".doc")

    def parse(self, file: InputFile) -> ParsedDocument:
        executable = shutil.which(self._command)
        if executable is None:
            raise DocumentParseError("解析 .doc 需要安装 LibreOffice，并确保 soffice 命令可用")
        with tempfile.TemporaryDirectory(prefix="shipping-billing-doc-") as directory:
            workdir = Path(directory)
            source = workdir / Path(file.name).name
            source.write_bytes(file.content)
            completed = subprocess.run(
                [
                    executable,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    directory,
                    str(source),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=60,
            )
            target = source.with_suffix(".docx")
            if completed.returncode != 0 or not target.exists():
                message = completed.stderr.strip() or completed.stdout.strip() or "未知转换错误"
                raise DocumentParseError(f"旧版 DOC 转换失败：{message}")
            parsed = self._docx_parser.parse(InputFile(target.name, target.read_bytes()))
            return ParsedDocument(
                source_name=file.name,
                source_type=parsed.source_type,
                text_blocks=parsed.text_blocks,
                tables=parsed.tables,
                raw_text=parsed.raw_text,
                warnings=("旧版 DOC 已通过 LibreOffice 转换为 DOCX 后解析",),
            )
