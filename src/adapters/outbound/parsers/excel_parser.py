"""XLS/XLSX workbook parsing."""

from io import BytesIO

import xlrd
from openpyxl import load_workbook

from adapters.outbound.parsers.base import DocumentParseError
from application.dto.documents import InputFile, ParsedDocument, SourceLocation, TableBlock


class ExcelParser:
    def supports(self, file: InputFile) -> bool:
        return file.name.lower().endswith((".xls", ".xlsx", ".xlsm"))

    def parse(self, file: InputFile) -> ParsedDocument:
        if file.name.lower().endswith(".xls"):
            return self._parse_xls(file)
        try:
            workbook = load_workbook(BytesIO(file.content), data_only=True, read_only=True)
        except Exception as exc:
            raise DocumentParseError(f"Excel 文件解析失败：{file.name}") from exc

        tables: list[TableBlock] = []
        raw_lines: list[str] = []
        for sheet in workbook.worksheets:
            rows = [tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)]
            non_empty = [row for row in rows if any(value not in (None, "") for value in row)]
            if not non_empty:
                continue
            headers = tuple(str(value or "") for value in non_empty[0])
            data_rows = tuple(tuple(value for value in row) for row in non_empty[1:])
            tables.append(
                TableBlock(
                    headers=headers,
                    rows=data_rows,
                    location=SourceLocation(source_name=file.name, sheet=sheet.title, row=1),
                )
            )
            raw_lines.append(f"[{sheet.title}]")
            raw_lines.extend(" | ".join(str(value or "") for value in row) for row in non_empty)
        return ParsedDocument(
            source_name=file.name,
            source_type="EXCEL",
            tables=tuple(tables),
            raw_text="\n".join(raw_lines),
        )

    def _parse_xls(self, file: InputFile) -> ParsedDocument:
        try:
            workbook = xlrd.open_workbook(file_contents=file.content)
        except Exception as exc:
            raise DocumentParseError(f"Excel 文件解析失败：{file.name}") from exc
        tables: list[TableBlock] = []
        raw_lines: list[str] = []
        for sheet in workbook.sheets():
            rows = [
                tuple(sheet.cell_value(row, col) for col in range(sheet.ncols))
                for row in range(sheet.nrows)
            ]
            non_empty = [row for row in rows if any(value not in (None, "") for value in row)]
            if not non_empty:
                continue
            headers = tuple(str(value or "") for value in non_empty[0])
            tables.append(
                TableBlock(
                    headers,
                    tuple(non_empty[1:]),
                    SourceLocation(file.name, sheet=sheet.name, row=1),
                )
            )
            raw_lines.append(f"[{sheet.name}]")
            raw_lines.extend(" | ".join(str(value or "") for value in row) for row in non_empty)
        return ParsedDocument(
            file.name, "EXCEL", tables=tuple(tables), raw_text="\n".join(raw_lines)
        )
